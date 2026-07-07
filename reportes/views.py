from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q, Max, Case, When, IntegerField
from .forms import ReporteForm
from .models import Reporte, Historial
from vecinos.models import Vecino
from django.utils import timezone
from datetime import timedelta
import csv



def inicio(request):
    """Landing institucional con indicadores reales del sistema."""
    total_reportes = Reporte.objects.count()
    reportes_resueltos = Reporte.objects.filter(estado='resuelto').count()
    reportes_en_gestion = Reporte.objects.exclude(estado='resuelto').count()
    reportes_con_gps = Reporte.objects.filter(latitud__isnull=False, longitud__isnull=False).count()

    porcentaje_resueltos = round((reportes_resueltos / total_reportes) * 100) if total_reportes else 0
    porcentaje_gps = round((reportes_con_gps / total_reportes) * 100) if total_reportes else 0

    context = {
        'total_reportes': total_reportes,
        'reportes_resueltos': reportes_resueltos,
        'reportes_en_gestion': reportes_en_gestion,
        'reportes_con_gps': reportes_con_gps,
        'porcentaje_resueltos': porcentaje_resueltos,
        'porcentaje_gps': porcentaje_gps,
    }
    return render(request, 'index.html', context)


def _historial_items(reporte):
    estados = dict(Reporte.ESTADOS)
    return [
        {
            'estado_anterior': h.estado_anterior,
            'estado_nuevo': h.estado_nuevo,
            'estado_anterior_display': estados.get(h.estado_anterior, h.estado_anterior),
            'estado_nuevo_display': estados.get(h.estado_nuevo, h.estado_nuevo),
            'fecha_cambio': h.fecha_cambio,
            'comentario': h.comentario,
        }
        for h in Historial.objects.filter(reporte=reporte).order_by('-fecha_cambio')
    ]

@login_required(login_url='/login/')
def reportar(request):
    if request.method == 'POST':
        form = ReporteForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                vecino = request.user.vecino
            except Vecino.DoesNotExist:
                messages.error(request, 'Tu usuario no tiene un perfil de vecino asociado.')
                return redirect('inicio')

            reporte = form.save(commit=False)
            reporte.vecino = vecino

            # Generar número de seguimiento único y estable
            ultimo_id = Reporte.objects.aggregate(Max('id'))['id__max'] or 0
            reporte.numero_seguimiento = f"REC-{ultimo_id + 1:04d}"

            reporte.save()
            messages.success(request, f'Reclamo {reporte.numero_seguimiento} registrado correctamente. Ya puedes hacer seguimiento desde esta ficha.')
            return redirect('detalle_reporte', numero_seguimiento=reporte.numero_seguimiento)
    else:
        form = ReporteForm()
    
    return render(request, 'reportar.html', {'form': form})


def consultar(request):
    reporte = None
    historial_items = []
    if request.method == 'POST':
        numero = request.POST.get('numero_seguimiento', '').strip().upper()
        try:
            reporte = Reporte.objects.get(numero_seguimiento=numero)
            historial_items = _historial_items(reporte)
        except Reporte.DoesNotExist:
            messages.error(request, 'No se encontró ningún reclamo con ese número de seguimiento')
    
    return render(request, 'consultar.html', {'reporte': reporte, 'historial_items': historial_items})


@login_required(login_url='/login/')
def dashboard(request):
    # Verificar que el usuario tenga perfil de vecino
    try:
        vecino = request.user.vecino
    except Vecino.DoesNotExist:
        messages.error(request, 'No tienes un perfil de vecino asociado.')
        return redirect('inicio')
    
    # Si es presidente, mostrar dashboard completo
    if vecino.rol == 'presidente':
        return dashboard_presidente(request, vecino)
    else:
        return dashboard_vecino(request, vecino)


def _filtrar_reportes_presidente(request, vecino):
    """Devuelve los reportes del barrio del presidente aplicando los filtros GET."""
    reportes_barrio = Reporte.objects.filter(vecino__barrio=vecino.barrio)

    filtro_estado = request.GET.get('estado', '').strip()
    filtro_tipo = request.GET.get('tipo', '').strip()
    filtro_busqueda = request.GET.get('q', '').strip()
    filtro_gps = request.GET.get('gps', '').strip()

    reportes_filtrados = reportes_barrio

    if filtro_estado:
        reportes_filtrados = reportes_filtrados.filter(estado=filtro_estado)

    if filtro_tipo:
        reportes_filtrados = reportes_filtrados.filter(tipo=filtro_tipo)

    if filtro_busqueda:
        reportes_filtrados = reportes_filtrados.filter(
            Q(numero_seguimiento__icontains=filtro_busqueda) |
            Q(ubicacion__icontains=filtro_busqueda) |
            Q(descripcion__icontains=filtro_busqueda) |
            Q(vecino__nombre__icontains=filtro_busqueda)
        )

    if filtro_gps == 'con_gps':
        reportes_filtrados = reportes_filtrados.filter(latitud__isnull=False, longitud__isnull=False)
    elif filtro_gps == 'sin_gps':
        reportes_filtrados = reportes_filtrados.filter(Q(latitud__isnull=True) | Q(longitud__isnull=True))

    filtros = {
        'estado': filtro_estado,
        'tipo': filtro_tipo,
        'q': filtro_busqueda,
        'gps': filtro_gps,
    }

    return reportes_barrio, reportes_filtrados, filtros


def dashboard_presidente(request, vecino):
    reportes_barrio, reportes_filtrados, filtros = _filtrar_reportes_presidente(request, vecino)

    # Estadísticas generales del barrio. No dependen de los filtros.
    total_reportes = reportes_barrio.count()
    reportes_pendientes = reportes_barrio.filter(~Q(estado='resuelto')).count()
    reportes_resueltos = reportes_barrio.filter(estado='resuelto').count()

    reportes_por_tipo = {
        'bache': reportes_barrio.filter(tipo='bache').count(),
        'basura': reportes_barrio.filter(tipo='basura').count(),
        'alumbrado': reportes_barrio.filter(tipo='alumbrado').count(),
        'otro': reportes_barrio.filter(tipo='otro').count(),
    }

    reportes_por_estado = {
        'reportado': reportes_barrio.filter(estado='reportado').count(),
        'validado': reportes_barrio.filter(estado='validado').count(),
        'derivado': reportes_barrio.filter(estado='derivado').count(),
        'proceso': reportes_barrio.filter(estado='proceso').count(),
        'resuelto': reportes_barrio.filter(estado='resuelto').count(),
    }

    cantidad_filtrada = reportes_filtrados.count()
    ultimos_reportes = reportes_filtrados.order_by('-fecha_creacion')[:50]

    # Indicadores premium para la gestión del barrio.
    reportes_con_gps = reportes_barrio.filter(latitud__isnull=False, longitud__isnull=False).count()
    reportes_sin_gps = total_reportes - reportes_con_gps
    porcentaje_resueltos = round((reportes_resueltos / total_reportes) * 100) if total_reportes else 0
    porcentaje_pendientes = round((reportes_pendientes / total_reportes) * 100) if total_reportes else 0
    porcentaje_gps = round((reportes_con_gps / total_reportes) * 100) if total_reportes else 0
    reportes_recientes_24h = reportes_barrio.filter(fecha_creacion__date__gte=timezone.localdate()).count()
    movimientos_recientes = Historial.objects.filter(
        reporte__vecino__barrio=vecino.barrio
    ).select_related('reporte').order_by('-fecha_cambio')[:6]

    # Centro de operaciones: alertas y prioridades de atención.
    limite_inactividad = timezone.now() - timedelta(days=7)
    reportes_nuevos = reportes_barrio.filter(estado='reportado').count()
    sin_actualizar_7d = reportes_barrio.filter(~Q(estado='resuelto'), fecha_actualizacion__lte=limite_inactividad).count()

    reportes_prioritarios = reportes_barrio.filter(~Q(estado='resuelto')).annotate(
        prioridad_orden=Case(
            When(estado='reportado', then=1),
            When(estado='validado', then=2),
            When(estado='derivado', then=3),
            When(estado='proceso', then=4),
            default=5,
            output_field=IntegerField(),
        )
    ).order_by('prioridad_orden', 'fecha_actualizacion')[:5]

    alertas_operativas = []
    if reportes_nuevos > 0:
        alertas_operativas.append({
            'icono': 'bi-bell',
            'titulo': 'Reclamos nuevos pendientes de validación',
            'detalle': f'{reportes_nuevos} reporte(s) todavía están en estado Reportado.',
            'tipo': 'warning',
        })
    if sin_actualizar_7d > 0:
        alertas_operativas.append({
            'icono': 'bi-clock-history',
            'titulo': 'Casos sin actualización reciente',
            'detalle': f'{sin_actualizar_7d} reclamo(s) llevan más de 7 días sin cambios.',
            'tipo': 'danger',
        })
    if reportes_sin_gps > 0:
        alertas_operativas.append({
            'icono': 'bi-geo-alt',
            'titulo': 'Reportes sin ubicación GPS',
            'detalle': f'{reportes_sin_gps} reporte(s) dependen solo de referencia escrita.',
            'tipo': 'info',
        })
    if reportes_pendientes == 0 and total_reportes > 0:
        alertas_operativas.append({
            'icono': 'bi-check2-circle',
            'titulo': 'Gestión al día',
            'detalle': 'Todos los reclamos del barrio están cerrados como resueltos.',
            'tipo': 'success',
        })
    if not alertas_operativas:
        alertas_operativas.append({
            'icono': 'bi-activity',
            'titulo': 'Operación estable',
            'detalle': 'No hay alertas críticas en este momento.',
            'tipo': 'success',
        })

    filtros_activos = any(filtros.values())

    context = {
        'vecino': vecino,
        'total_reportes': total_reportes,
        'reportes_pendientes': reportes_pendientes,
        'reportes_resueltos': reportes_resueltos,
        'reportes_por_tipo': reportes_por_tipo,
        'reportes_por_estado': reportes_por_estado,
        'ultimos_reportes': ultimos_reportes,
        'cantidad_filtrada': cantidad_filtrada,
        'reportes_con_gps': reportes_con_gps,
        'reportes_sin_gps': reportes_sin_gps,
        'porcentaje_resueltos': porcentaje_resueltos,
        'porcentaje_pendientes': porcentaje_pendientes,
        'porcentaje_gps': porcentaje_gps,
        'reportes_recientes_24h': reportes_recientes_24h,
        'movimientos_recientes': movimientos_recientes,
        'alertas_operativas': alertas_operativas,
        'reportes_prioritarios': reportes_prioritarios,
        'reportes_nuevos': reportes_nuevos,
        'sin_actualizar_7d': sin_actualizar_7d,
        'filtros': filtros,
        'filtros_activos': filtros_activos,
        'tipos_reporte': Reporte.TIPOS,
        'estados_reporte': Reporte.ESTADOS,
        'es_presidente': True,
    }

    return render(request, 'dashboard_presidente.html', context)


@login_required(login_url='/login/')
def centro_reportes(request):
    """Centro institucional de reportes para presidentes de junta."""
    try:
        vecino = request.user.vecino
    except Vecino.DoesNotExist:
        messages.error(request, 'No tienes un perfil de vecino asociado.')
        return redirect('inicio')

    if vecino.rol != 'presidente':
        messages.error(request, 'No tienes permiso para acceder al Centro de Reportes.')
        return redirect('dashboard')

    reportes_barrio, reportes_filtrados, filtros = _filtrar_reportes_presidente(request, vecino)
    reportes_filtrados = reportes_filtrados.select_related('vecino').order_by('-fecha_creacion')

    total_general = reportes_barrio.count()
    total_filtrado = reportes_filtrados.count()
    resueltos_filtrados = reportes_filtrados.filter(estado='resuelto').count()
    pendientes_filtrados = reportes_filtrados.filter(~Q(estado='resuelto')).count()
    con_gps_filtrados = reportes_filtrados.filter(latitud__isnull=False, longitud__isnull=False).count()
    porcentaje_resueltos_filtrado = round((resueltos_filtrados / total_filtrado) * 100) if total_filtrado else 0
    porcentaje_gps_filtrado = round((con_gps_filtrados / total_filtrado) * 100) if total_filtrado else 0

    resumen_por_estado = reportes_filtrados.values('estado').annotate(total=Count('id')).order_by('estado')
    resumen_por_tipo = reportes_filtrados.values('tipo').annotate(total=Count('id')).order_by('tipo')

    context = {
        'vecino': vecino,
        'filtros': filtros,
        'filtros_activos': any(filtros.values()),
        'tipos_reporte': Reporte.TIPOS,
        'estados_reporte': Reporte.ESTADOS,
        'total_general': total_general,
        'total_filtrado': total_filtrado,
        'resueltos_filtrados': resueltos_filtrados,
        'pendientes_filtrados': pendientes_filtrados,
        'con_gps_filtrados': con_gps_filtrados,
        'porcentaje_resueltos_filtrado': porcentaje_resueltos_filtrado,
        'porcentaje_gps_filtrado': porcentaje_gps_filtrado,
        'resumen_por_estado': resumen_por_estado,
        'resumen_por_tipo': resumen_por_tipo,
        'reportes_preview': reportes_filtrados[:10],
    }
    return render(request, 'centro_reportes.html', context)


@login_required(login_url='/login/')
def exportar_reportes_csv(request):
    """Exporta los reportes del barrio del presidente respetando los filtros actuales."""
    try:
        vecino = request.user.vecino
    except Vecino.DoesNotExist:
        messages.error(request, 'No tienes un perfil de vecino asociado.')
        return redirect('inicio')

    if vecino.rol != 'presidente':
        messages.error(request, 'No tienes permiso para exportar reportes.')
        return redirect('dashboard')

    _, reportes_filtrados, _ = _filtrar_reportes_presidente(request, vecino)
    reportes_filtrados = reportes_filtrados.select_related('vecino').order_by('-fecha_creacion')

    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    fecha_archivo = timezone.localtime(timezone.now()).strftime('%Y-%m-%d_%H-%M')
    response['Content-Disposition'] = f'attachment; filename="SIGEREC_reportes_{fecha_archivo}.csv"'
    response.write('\ufeff')

    # Separador punto y coma para mejor compatibilidad con Excel en español.
    writer = csv.writer(response, delimiter=';')
    writer.writerow([
        'Numero de seguimiento',
        'Vecino',
        'CI',
        'Telefono',
        'Barrio',
        'Tipo',
        'Estado',
        'Ubicacion escrita',
        'Latitud',
        'Longitud',
        'Descripcion',
        'Fecha de creacion',
        'Ultima actualizacion',
    ])

    for reporte in reportes_filtrados:
        writer.writerow([
            reporte.numero_seguimiento,
            reporte.vecino.nombre,
            reporte.vecino.ci,
            reporte.vecino.telefono,
            reporte.vecino.barrio,
            reporte.get_tipo_display(),
            reporte.get_estado_display(),
            reporte.ubicacion,
            reporte.latitud or '',
            reporte.longitud or '',
            reporte.descripcion,
            timezone.localtime(reporte.fecha_creacion).strftime('%d/%m/%Y %H:%M'),
            timezone.localtime(reporte.fecha_actualizacion).strftime('%d/%m/%Y %H:%M'),
        ])

    return response



@login_required(login_url='/login/')
def exportar_reportes_excel(request):
    """Genera un informe institucional en Excel con resumen, filtros y detalle de reclamos."""
    try:
        vecino = request.user.vecino
    except Vecino.DoesNotExist:
        messages.error(request, 'No tienes un perfil de vecino asociado.')
        return redirect('inicio')

    if vecino.rol != 'presidente':
        messages.error(request, 'No tienes permiso para exportar reportes.')
        return redirect('dashboard')

    try:
        from io import BytesIO
        from pathlib import Path
        from django.conf import settings
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        from openpyxl.utils import get_column_letter
        from openpyxl.drawing.image import Image as ExcelImage
    except ImportError:
        messages.error(request, 'Falta instalar openpyxl. Ejecuta: pip install openpyxl')
        return redirect('centro_reportes')

    _, reportes_filtrados, filtros = _filtrar_reportes_presidente(request, vecino)
    reportes_filtrados = reportes_filtrados.select_related('vecino').order_by('-fecha_creacion')
    reportes = list(reportes_filtrados)

    total = len(reportes)
    pendientes = sum(1 for r in reportes if r.estado != 'resuelto')
    resueltos = sum(1 for r in reportes if r.estado == 'resuelto')
    con_gps = sum(1 for r in reportes if r.latitud is not None and r.longitud is not None)
    sin_gps = total - con_gps
    porcentaje_resueltos = round((resueltos / total) * 100) if total else 0
    porcentaje_gps = round((con_gps / total) * 100) if total else 0

    wb = Workbook()
    ws = wb.active
    ws.title = 'Informe SIGEREC'
    detalle = wb.create_sheet('Detalle de reclamos')

    # Paleta institucional
    verde = '123D32'
    verde_claro = '2E7D5B'
    azul = '2D79C7'
    gris_fondo = 'F4F7FA'
    gris_borde = 'D9E2EC'
    blanco = 'FFFFFF'
    texto = '183247'
    amarillo = 'FFF4CC'
    rojo = 'FFE1E1'
    celeste = 'E8F2FF'
    verde_suave = 'DDF7EA'

    thin = Side(style='thin', color=gris_borde)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet in (ws, detalle):
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = 'A12' if sheet == ws else 'A2'

    # Portada / encabezado
    ws.merge_cells('A1:M1')
    ws['A1'] = 'GOBIERNO AUTÓNOMO MUNICIPAL DE TRINIDAD'
    ws['A1'].font = Font(bold=True, size=16, color=blanco)
    ws['A1'].alignment = Alignment(horizontal='center')
    ws['A1'].fill = PatternFill('solid', fgColor=verde)

    ws.merge_cells('A2:M2')
    ws['A2'] = 'SIGEREC — Sistema de Gestión de Reclamos Ciudadanos'
    ws['A2'].font = Font(bold=True, size=14, color=blanco)
    ws['A2'].alignment = Alignment(horizontal='center')
    ws['A2'].fill = PatternFill('solid', fgColor=verde_claro)

    ws.merge_cells('A3:M3')
    ws['A3'] = 'Informe institucional generado automáticamente'
    ws['A3'].font = Font(italic=True, size=11, color=texto)
    ws['A3'].alignment = Alignment(horizontal='center')
    ws['A3'].fill = PatternFill('solid', fgColor=gris_fondo)

    logo_path = Path(settings.BASE_DIR) / 'static' / 'images' / 'logo.png'
    if logo_path.exists():
        try:
            logo = ExcelImage(str(logo_path))
            logo.height = 52
            logo.width = 135
            ws.add_image(logo, 'A1')
        except Exception:
            pass

    ahora = timezone.localtime(timezone.now())
    datos_generales = [
        ('Fecha de generación', ahora.strftime('%d/%m/%Y')),
        ('Hora de generación', ahora.strftime('%H:%M')),
        ('Usuario responsable', vecino.nombre),
        ('Barrio gestionado', vecino.barrio),
        ('Cantidad exportada', total),
    ]

    ws['A5'] = 'Datos del informe'
    ws['A5'].font = Font(bold=True, size=13, color=verde)
    for idx, (clave, valor) in enumerate(datos_generales, start=6):
        ws[f'A{idx}'] = clave
        ws[f'B{idx}'] = valor
        ws[f'A{idx}'].font = Font(bold=True, color=texto)
        ws[f'A{idx}'].fill = PatternFill('solid', fgColor=gris_fondo)
        ws[f'A{idx}'].border = border
        ws[f'B{idx}'].border = border

    filtros_legibles = []
    if filtros.get('q'):
        filtros_legibles.append(('Búsqueda', filtros['q']))
    if filtros.get('estado'):
        filtros_legibles.append(('Estado', dict(Reporte.ESTADOS).get(filtros['estado'], filtros['estado'])))
    if filtros.get('tipo'):
        filtros_legibles.append(('Tipo', dict(Reporte.TIPOS).get(filtros['tipo'], filtros['tipo'])))
    if filtros.get('gps'):
        filtros_legibles.append(('GPS', 'Solo con GPS' if filtros['gps'] == 'con_gps' else 'Solo sin GPS'))
    if not filtros_legibles:
        filtros_legibles.append(('Filtros', 'Sin filtros aplicados'))

    ws['D5'] = 'Filtros aplicados'
    ws['D5'].font = Font(bold=True, size=13, color=verde)
    for idx, (clave, valor) in enumerate(filtros_legibles, start=6):
        ws[f'D{idx}'] = clave
        ws[f'E{idx}'] = valor
        ws[f'D{idx}'].font = Font(bold=True, color=texto)
        ws[f'D{idx}'].fill = PatternFill('solid', fgColor=gris_fondo)
        ws[f'D{idx}'].border = border
        ws[f'E{idx}'].border = border

    resumen = [
        ('Total de reclamos', total, celeste),
        ('Pendientes', pendientes, amarillo),
        ('Resueltos', resueltos, verde_suave),
        ('Con GPS', con_gps, celeste),
        ('Sin GPS', sin_gps, rojo if sin_gps else verde_suave),
        ('% Resueltos', f'{porcentaje_resueltos}%', verde_suave),
        ('% con GPS', f'{porcentaje_gps}%', celeste),
    ]

    ws['G5'] = 'Resumen ejecutivo'
    ws['G5'].font = Font(bold=True, size=13, color=verde)
    for idx, (clave, valor, color) in enumerate(resumen, start=6):
        ws[f'G{idx}'] = clave
        ws[f'H{idx}'] = valor
        ws[f'G{idx}'].font = Font(bold=True, color=texto)
        ws[f'H{idx}'].font = Font(bold=True, size=12, color=texto)
        ws[f'G{idx}'].fill = PatternFill('solid', fgColor=gris_fondo)
        ws[f'H{idx}'].fill = PatternFill('solid', fgColor=color)
        ws[f'G{idx}'].border = border
        ws[f'H{idx}'].border = border

    encabezados = [
        'N°', 'Código', 'Vecino', 'CI', 'Teléfono', 'Barrio', 'Tipo', 'Estado',
        'Ubicación escrita', 'GPS', 'Latitud', 'Longitud', 'Google Maps',
        'Descripción', 'Fecha de creación', 'Última actualización'
    ]

    start_row = 14
    for col, header in enumerate(encabezados, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font = Font(bold=True, color=blanco)
        cell.fill = PatternFill('solid', fgColor=verde)
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    fill_estado = {
        'reportado': rojo,
        'validado': celeste,
        'derivado': 'EFE6FF',
        'proceso': amarillo,
        'resuelto': verde_suave,
    }

    for row_idx, reporte in enumerate(reportes, start=start_row + 1):
        maps_url = ''
        tiene_gps = reporte.latitud is not None and reporte.longitud is not None
        if tiene_gps:
            maps_url = f'https://www.google.com/maps?q={reporte.latitud},{reporte.longitud}'

        valores = [
            row_idx - start_row,
            reporte.numero_seguimiento,
            reporte.vecino.nombre,
            reporte.vecino.ci,
            reporte.vecino.telefono,
            reporte.vecino.barrio,
            reporte.get_tipo_display(),
            reporte.get_estado_display(),
            reporte.ubicacion,
            'Sí' if tiene_gps else 'No',
            float(reporte.latitud) if reporte.latitud is not None else '',
            float(reporte.longitud) if reporte.longitud is not None else '',
            'Abrir mapa' if maps_url else '',
            reporte.descripcion,
            timezone.localtime(reporte.fecha_creacion).strftime('%d/%m/%Y %H:%M'),
            timezone.localtime(reporte.fecha_actualizacion).strftime('%d/%m/%Y %H:%M'),
        ]

        for col, value in enumerate(valores, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if row_idx % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='FAFCFE')

        estado_cell = ws.cell(row=row_idx, column=8)
        estado_cell.fill = PatternFill('solid', fgColor=fill_estado.get(reporte.estado, gris_fondo))
        estado_cell.font = Font(bold=True, color=texto)

        gps_cell = ws.cell(row=row_idx, column=10)
        gps_cell.fill = PatternFill('solid', fgColor=verde_suave if tiene_gps else rojo)
        gps_cell.font = Font(bold=True, color=texto)
        gps_cell.alignment = Alignment(horizontal='center')

        if maps_url:
            map_cell = ws.cell(row=row_idx, column=13)
            map_cell.hyperlink = maps_url
            map_cell.style = 'Hyperlink'

    ws.auto_filter.ref = f'A{start_row}:P{start_row + max(total, 1)}'
    ws.freeze_panes = f'A{start_row + 1}'

    ws.merge_cells(start_row=start_row + total + 3, start_column=1, end_row=start_row + total + 3, end_column=16)
    footer = ws.cell(row=start_row + total + 3, column=1)
    footer.value = 'Documento generado automáticamente por SIGEREC — Proyecto de Extensión Universitaria — Universidad Autónoma del Beni'
    footer.font = Font(italic=True, color='6B7B8C')
    footer.alignment = Alignment(horizontal='center')

    # Hoja de detalle limpia para análisis de datos.
    for col, header in enumerate(encabezados, start=1):
        cell = detalle.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color=blanco)
        cell.fill = PatternFill('solid', fgColor=verde)
        cell.border = border
        cell.alignment = Alignment(horizontal='center')

    for row_idx, reporte in enumerate(reportes, start=2):
        tiene_gps = reporte.latitud is not None and reporte.longitud is not None
        maps_url = f'https://www.google.com/maps?q={reporte.latitud},{reporte.longitud}' if tiene_gps else ''
        valores = [
            row_idx - 1,
            reporte.numero_seguimiento,
            reporte.vecino.nombre,
            reporte.vecino.ci,
            reporte.vecino.telefono,
            reporte.vecino.barrio,
            reporte.get_tipo_display(),
            reporte.get_estado_display(),
            reporte.ubicacion,
            'Sí' if tiene_gps else 'No',
            float(reporte.latitud) if reporte.latitud is not None else '',
            float(reporte.longitud) if reporte.longitud is not None else '',
            maps_url,
            reporte.descripcion,
            timezone.localtime(reporte.fecha_creacion).strftime('%d/%m/%Y %H:%M'),
            timezone.localtime(reporte.fecha_actualizacion).strftime('%d/%m/%Y %H:%M'),
        ]
        for col, value in enumerate(valores, start=1):
            cell = detalle.cell(row=row_idx, column=col, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if col == 13 and maps_url:
                cell.hyperlink = maps_url
                cell.style = 'Hyperlink'

    detalle.auto_filter.ref = f'A1:P{max(total + 1, 2)}'

    widths = {
        1: 6, 2: 14, 3: 26, 4: 14, 5: 14, 6: 20, 7: 22, 8: 22,
        9: 36, 10: 10, 11: 14, 12: 14, 13: 18, 14: 45, 15: 20, 16: 20,
    }
    for sheet in (ws, detalle):
        for col_idx, width in widths.items():
            sheet.column_dimensions[get_column_letter(col_idx)].width = width
        for row in range(1, sheet.max_row + 1):
            sheet.row_dimensions[row].height = 24

    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[14].height = 34
    detalle.row_dimensions[1].height = 32

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    fecha_archivo = ahora.strftime('%Y-%m-%d_%H-%M')
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="SIGEREC_informe_institucional_{fecha_archivo}.xlsx"'
    return response

def dashboard_vecino(request, vecino):
    # Obtener solo los reportes del vecino
    reportes = Reporte.objects.filter(vecino=vecino)
    
    # Estadísticas personales
    total_reportes = reportes.count()
    reportes_pendientes = reportes.filter(~Q(estado='resuelto')).count()
    reportes_resueltos = reportes.filter(estado='resuelto').count()
    
    # Reportes por tipo (personales)
    reportes_por_tipo = {
        'bache': reportes.filter(tipo='bache').count(),
        'basura': reportes.filter(tipo='basura').count(),
        'alumbrado': reportes.filter(tipo='alumbrado').count(),
        'otro': reportes.filter(tipo='otro').count(),
    }
    
    # Reportes por estado (personales)
    reportes_por_estado = {
        'reportado': reportes.filter(estado='reportado').count(),
        'validado': reportes.filter(estado='validado').count(),
        'derivado': reportes.filter(estado='derivado').count(),
        'proceso': reportes.filter(estado='proceso').count(),
        'resuelto': reportes.filter(estado='resuelto').count(),
    }
    
    # Últimos 5 reportes del vecino
    ultimos_reportes = reportes.order_by('-fecha_creacion')[:5]
    
    context = {
        'vecino': vecino,
        'total_reportes': total_reportes,
        'reportes_pendientes': reportes_pendientes,
        'reportes_resueltos': reportes_resueltos,
        'reportes_por_tipo': reportes_por_tipo,
        'reportes_por_estado': reportes_por_estado,
        'ultimos_reportes': ultimos_reportes,
        'es_presidente': False,
    }
    
    return render(request, 'dashboard_vecino.html', context)


def detalle_reporte(request, numero_seguimiento):
    numero = numero_seguimiento.strip().upper()
    reporte = get_object_or_404(Reporte, numero_seguimiento=numero)
    historial_items = _historial_items(reporte)
    return render(request, 'detalle_reporte.html', {
        'reporte': reporte,
        'historial_items': historial_items,
    })


@login_required(login_url='/login/')
def cambiar_estado(request, reporte_id):
    # Verificar que el usuario sea presidente
    try:
        vecino = request.user.vecino
        if vecino.rol != 'presidente':
            messages.error(request, 'No tienes permiso para realizar esta acción')
            return redirect('dashboard')
    except Vecino.DoesNotExist:
        messages.error(request, 'No tienes permiso para realizar esta acción')
        return redirect('dashboard')
    
    reporte = get_object_or_404(Reporte, id=reporte_id, vecino__barrio=vecino.barrio)
    
    if request.method == 'POST':
        nuevo_estado = request.POST.get('estado')
        comentario = request.POST.get('comentario', '')
        
        if nuevo_estado in dict(Reporte.ESTADOS):
            # Guardar historial
            Historial.objects.create(
                reporte=reporte,
                estado_anterior=reporte.estado,
                estado_nuevo=nuevo_estado,
                comentario=comentario
            )
            
            # Actualizar estado
            reporte.estado = nuevo_estado
            reporte.save()
            
            messages.success(request, f'Estado actualizado a: {dict(Reporte.ESTADOS)[nuevo_estado]}')
        else:
            messages.error(request, 'Estado no válido')
    
    return redirect('dashboard')


@login_required(login_url='/login/')
def exportar_reportes_pdf(request):
    """Genera un informe institucional PDF de reclamos del barrio del presidente."""
    try:
        vecino = request.user.vecino
    except Vecino.DoesNotExist:
        messages.error(request, 'No tienes un perfil de vecino asociado.')
        return redirect('inicio')

    if vecino.rol != 'presidente':
        messages.error(request, 'No tienes permiso para exportar reportes.')
        return redirect('dashboard')

    try:
        from io import BytesIO
        from pathlib import Path
        from django.conf import settings
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
            Image, PageBreak
        )
    except ImportError:
        messages.error(request, 'Falta instalar ReportLab. Ejecuta: pip install reportlab')
        return redirect('centro_reportes')

    _, reportes_filtrados, filtros = _filtrar_reportes_presidente(request, vecino)
    reportes = list(reportes_filtrados.select_related('vecino').order_by('-fecha_creacion'))

    total = len(reportes)
    pendientes = sum(1 for r in reportes if r.estado != 'resuelto')
    resueltos = sum(1 for r in reportes if r.estado == 'resuelto')
    con_gps = sum(1 for r in reportes if r.latitud is not None and r.longitud is not None)
    porcentaje_resueltos = round((resueltos / total) * 100) if total else 0
    porcentaje_gps = round((con_gps / total) * 100) if total else 0

    buffer = BytesIO()
    fecha_archivo = timezone.localtime(timezone.now()).strftime('%Y-%m-%d_%H-%M')
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=1.2 * cm,
        leftMargin=1.2 * cm,
        topMargin=1.0 * cm,
        bottomMargin=1.0 * cm,
        title='Informe Institucional SIGEREC',
        author='SIGEREC',
    )

    styles = getSampleStyleSheet()
    primary = colors.HexColor('#123D32')
    secondary = colors.HexColor('#2E7D5B')
    accent = colors.HexColor('#2D79C7')
    soft_bg = colors.HexColor('#F4F7FA')
    border = colors.HexColor('#D9E2EC')
    text_color = colors.HexColor('#183247')

    title_style = ParagraphStyle(
        'SIGERECTitle', parent=styles['Title'], alignment=TA_CENTER,
        fontSize=18, leading=22, textColor=primary, spaceAfter=8
    )
    subtitle_style = ParagraphStyle(
        'SIGERECSubtitle', parent=styles['Normal'], alignment=TA_CENTER,
        fontSize=10, leading=14, textColor=text_color, spaceAfter=12
    )
    section_style = ParagraphStyle(
        'SIGERECSection', parent=styles['Heading2'], fontSize=12,
        leading=15, textColor=primary, spaceBefore=10, spaceAfter=6
    )
    normal_style = ParagraphStyle(
        'SIGERECNormal', parent=styles['Normal'], fontSize=8.2,
        leading=10.5, textColor=text_color
    )
    small_style = ParagraphStyle(
        'SIGERECSmall', parent=styles['Normal'], fontSize=7.2,
        leading=9, textColor=text_color
    )

    elements = []

    logo_path = Path(settings.BASE_DIR) / 'static' / 'images' / 'logo.png'
    if logo_path.exists():
        try:
            logo = Image(str(logo_path), width=3.0*cm, height=1.1*cm)
            logo.hAlign = 'CENTER'
            elements.append(logo)
            elements.append(Spacer(1, 0.15*cm))
        except Exception:
            pass

    elements.append(Paragraph('GOBIERNO AUTÓNOMO MUNICIPAL DE TRINIDAD', title_style))
    elements.append(Paragraph('SIGEREC — Sistema de Gestión de Reclamos Ciudadanos', subtitle_style))
    elements.append(Paragraph('Informe Institucional de Reclamos Ciudadanos', title_style))

    ahora = timezone.localtime(timezone.now())
    datos = [
        ['Fecha de generación', ahora.strftime('%d/%m/%Y')],
        ['Hora de generación', ahora.strftime('%H:%M')],
        ['Responsable', vecino.nombre],
        ['Barrio gestionado', vecino.barrio],
    ]
    filtros_legibles = []
    if filtros.get('q'):
        filtros_legibles.append(['Búsqueda', filtros['q']])
    if filtros.get('estado'):
        filtros_legibles.append(['Estado', dict(Reporte.ESTADOS).get(filtros['estado'], filtros['estado'])])
    if filtros.get('tipo'):
        filtros_legibles.append(['Tipo', dict(Reporte.TIPOS).get(filtros['tipo'], filtros['tipo'])])
    if filtros.get('gps'):
        filtros_legibles.append(['GPS', 'Solo con GPS' if filtros['gps'] == 'con_gps' else 'Solo sin GPS'])
    if not filtros_legibles:
        filtros_legibles.append(['Filtros', 'Sin filtros aplicados'])

    info_table = Table(
        [[Paragraph('<b>Datos del informe</b>', normal_style), '', Paragraph('<b>Filtros aplicados</b>', normal_style), '']] +
        [[Paragraph(k, small_style), Paragraph(str(v), small_style), Paragraph(filtros_legibles[i][0], small_style) if i < len(filtros_legibles) else '', Paragraph(str(filtros_legibles[i][1]), small_style) if i < len(filtros_legibles) else ''] for i, (k, v) in enumerate(datos)],
        colWidths=[4.0*cm, 6.0*cm, 4.0*cm, 10.0*cm]
    )
    info_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), soft_bg),
        ('TEXTCOLOR', (0,0), (-1,-1), text_color),
        ('GRID', (0,0), (-1,-1), 0.35, border),
        ('SPAN', (0,0), (1,0)),
        ('SPAN', (2,0), (3,0)),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('LEFTPADDING', (0,0), (-1,-1), 6),
        ('RIGHTPADDING', (0,0), (-1,-1), 6),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
    ]))
    elements.append(info_table)
    elements.append(Spacer(1, 0.35*cm))

    elements.append(Paragraph('Resumen ejecutivo', section_style))
    summary_data = [
        ['Total', 'Pendientes', 'Resueltos', 'Con GPS', '% Resueltos', '% GPS'],
        [str(total), str(pendientes), str(resueltos), str(con_gps), f'{porcentaje_resueltos}%', f'{porcentaje_gps}%'],
    ]
    summary_table = Table(summary_data, colWidths=[4.1*cm]*6)
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), primary),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('BACKGROUND', (0,1), (-1,1), soft_bg),
        ('TEXTCOLOR', (0,1), (-1,1), text_color),
        ('FONTNAME', (0,0), (-1,-1), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,0), 8),
        ('FONTSIZE', (0,1), (-1,1), 14),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('GRID', (0,0), (-1,-1), 0.35, border),
        ('TOPPADDING', (0,0), (-1,-1), 7),
        ('BOTTOMPADDING', (0,0), (-1,-1), 7),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 0.35*cm))

    elements.append(Paragraph('Detalle de reclamos', section_style))
    table_data = [[
        'Código', 'Vecino', 'Barrio', 'Tipo', 'Estado', 'Ubicación', 'GPS', 'Fecha'
    ]]

    for r in reportes:
        gps = 'Sí' if r.latitud is not None and r.longitud is not None else 'No'
        table_data.append([
            Paragraph(r.numero_seguimiento, small_style),
            Paragraph(r.vecino.nombre, small_style),
            Paragraph(r.vecino.barrio, small_style),
            Paragraph(r.get_tipo_display(), small_style),
            Paragraph(r.get_estado_display(), small_style),
            Paragraph((r.ubicacion or '')[:85], small_style),
            Paragraph(gps, small_style),
            Paragraph(timezone.localtime(r.fecha_creacion).strftime('%d/%m/%Y'), small_style),
        ])

    if len(table_data) == 1:
        table_data.append([Paragraph('No existen reclamos para los filtros aplicados.', normal_style), '', '', '', '', '', '', ''])

    detail_table = Table(table_data, repeatRows=1, colWidths=[2.7*cm, 4.0*cm, 3.3*cm, 2.7*cm, 2.8*cm, 7.2*cm, 1.3*cm, 2.3*cm])
    style_cmds = [
        ('BACKGROUND', (0,0), (-1,0), secondary),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,0), 7.5),
        ('ALIGN', (0,0), (-1,0), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('GRID', (0,0), (-1,-1), 0.25, border),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, soft_bg]),
        ('LEFTPADDING', (0,0), (-1,-1), 4),
        ('RIGHTPADDING', (0,0), (-1,-1), 4),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]
    if len(table_data) == 2 and total == 0:
        style_cmds.append(('SPAN', (0,1), (-1,1)))
        style_cmds.append(('ALIGN', (0,1), (-1,1), 'CENTER'))
    detail_table.setStyle(TableStyle(style_cmds))
    elements.append(detail_table)

    elements.append(Spacer(1, 0.4*cm))
    elements.append(Paragraph('Documento generado automáticamente por SIGEREC. Proyecto de extensión universitaria — Universidad Autónoma del Beni “José Ballivián”.', subtitle_style))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 7)
        canvas.setFillColor(text_color)
        canvas.drawString(1.2*cm, 0.55*cm, 'SIGEREC · Sistema de Gestión de Reclamos Ciudadanos')
        canvas.drawRightString(28.5*cm, 0.55*cm, f'Página {doc.page}')
        canvas.restoreState()

    doc.build(elements, onFirstPage=footer, onLaterPages=footer)
    buffer.seek(0)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="SIGEREC_informe_{fecha_archivo}.pdf"'
    return response
